"""Exportação do plano para Excel e PDF (bytes, prontos para o botão de download do painel)."""
from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pandas as pd

from .historico import CAMPOS_COMPARAR
from .regras import NIVEIS, ROT, br, br_kg

NOMES_XLSX = {
    "sigla": "Sigla", "cultura_rotulo": "Cultura", "area_m2": "Área (m²)", "ssp_kg_ha": "Superfosfato (kg/ha)",
    "ssp_a_cada_anos": "Superfosfato a cada (anos)", "ssp_kg": "Superfosfato no talhão (kg)",
    "ssp_g_m2": "Superfosfato (g/m²)", "npk_kg_ha": "25-00-25 (kg/ha/ano)", "npk_kg": "25-00-25 no talhão (kg)",
    "npk_g_m2": "25-00-25 (g/m² no ano)", "parcelas": "Parcelas", "quando_npk": "Quando (25-00-25)",
    "kcl_kg_ha": "Cloreto de potássio (kg/ha)", "kcl_kg": "Cloreto de potássio no talhão (kg)",
    "calcario_t_ha": "Calcário (t/ha)", "calcario_tipo": "Tipo de calcário", "calcario_kg": "Calcário no talhão (kg)",
    "N_kg_ha": "N (kg/ha)", "P2O5_kg_ha": "P2O5 (kg/ha)", "K2O_kg_ha": "K2O (kg/ha)", "nota": "Como aplicar",
}


def gerar_xlsx(res: dict, laudo: pd.DataFrame, comparacao: pd.DataFrame | None = None,
               alertas_historico: list[str] | None = None) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill

    plano = res["plano"]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        cols = [c for c in NOMES_XLSX if c in plano]
        plano[cols].rename(columns=NOMES_XLSX).round(2).to_excel(xw, sheet_name="Plano por talhão", index=False)
        compras = res["compras"].rename(columns={"produto": "Produto", "kg": "Total (kg)", "embalagens": "Embalagens",
                                                 "embalagem": "Embalagem", "preco_embalagem": "Preço da embalagem (R$)",
                                                 "custo": "Custo (R$)"})
        compras.assign(Produto=compras["Produto"].str.replace("P₂O₅", "P2O5")).to_excel(
            xw, sheet_name="Compras", index=False)
        res["calendario"].rename(columns={"sigla": "Sigla"}).to_excel(xw, sheet_name="Calendário", index=False)
        if not res["coqueiros"].empty:
            res["coqueiros"].rename(columns={
                "sigla": "Sigla", "n_plantas": "Coqueiros", "idade": "Idade (anos)", "npk_kg": "25-00-25 (kg/planta/ano)",
                "ssp_kg": "Superfosfato (kg/planta)", "borax_g": "Bórax (g/planta)", "parcelas": "Parcelas",
                "kg_parcela": "25-00-25 por parcela (kg)", "npk_total_kg": "25-00-25 total (kg)",
                "ssp_total_kg": "Superfosfato total (kg)", "borax_total_kg": "Bórax total (kg)"}
            ).to_excel(xw, sheet_name="Coqueiros", index=False)
        laudo.to_excel(xw, sheet_name="Laudo", index=False)
        if comparacao is not None and not comparacao.empty:
            comparacao.to_excel(xw, sheet_name="Comparação", index=False)
        alertas = [*(alertas_historico or []), *res["alertas"]] or ["—"]
        pd.DataFrame({"Alertas e observações": alertas}).to_excel(xw, sheet_name="Alertas", index=False)

        verde, amarelo = PatternFill("solid", fgColor="2E5E3A"), PatternFill("solid", fgColor="FFF2CC")
        for ws in xw.book.worksheets:
            for cel in ws[1]:
                cel.fill, cel.font = verde, Font(bold=True, color="FFFFFF")
                cel.alignment = Alignment(wrap_text=True, vertical="center")
            for col in ws.columns:
                largura = max(len(str(c.value or "")) for c in col[:50])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, largura * 0.9), 45)
        ws = xw.book["Compras"]
        for i in range(2, ws.max_row + 1):  # preço editável, custo por fórmula
            ws[f"E{i}"].fill = amarelo
            ws[f"F{i}"] = f'=IF(E{i}="","",C{i}*E{i})'
        n = ws.max_row
        ws[f"E{n + 1}"], ws[f"F{n + 1}"] = "Total", f"=SUM(F2:F{n})"
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
def _formula(texto: str) -> str:
    """P₂O₅ com <sub> do reportlab: a Arial não tem os glifos ₂ e ₅ (a DejaVu tem; a marcação serve às duas)."""
    return texto.replace("P₂O₅", "P<sub>2</sub>O<sub>5</sub>")


def _estilos():
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    # Arial do Windows; no navegador (Pyodide), a DejaVu Sans empacotada em solo/fontes. As duas têm "→" e "≤",
    # que a Helvetica embutida do reportlab não tem.
    fontes = Path(__file__).parent / "fontes"
    F, FB = "Helvetica", "Helvetica-Bold"
    for familia, pasta, arqs in (
            ("Arial", Path("C:/Windows/Fonts"), ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf")),
            # o relatório não usa itálico: a versão online leva só a regular e a negrito
            ("DejaVuSans", fontes, ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans.ttf",
                                    "DejaVuSans-Bold.ttf"))):
        if not (pasta / arqs[0]).is_file():  # no navegador, "C:/..." viraria um pedido de URL
            continue
        try:
            nomes = [familia, f"{familia}-Bold", f"{familia}-Italic", f"{familia}-BoldItalic"]
            for nome, arq in zip(nomes, arqs):
                pdfmetrics.registerFont(TTFont(nome, str(pasta / arq)))
            pdfmetrics.registerFontFamily(familia, normal=nomes[0], bold=nomes[1], italic=nomes[2],
                                          boldItalic=nomes[3])
            F, FB = nomes[0], nomes[1]
            break
        except Exception:
            continue
    verde, cinza = colors.HexColor("#2E5E3A"), colors.HexColor("#5F6368")
    S = {
        "titulo": ParagraphStyle("t", fontName=FB, fontSize=16, leading=20, textColor=verde),
        "sub": ParagraphStyle("s", fontName=F, fontSize=8.8, leading=12, textColor=cinza),
        "h1": ParagraphStyle("h1", fontName=FB, fontSize=11.5, leading=15, textColor=verde, spaceBefore=10, spaceAfter=4),
        "corpo": ParagraphStyle("c", fontName=F, fontSize=9, leading=12.6, spaceAfter=3),
        "bullet": ParagraphStyle("b", fontName=F, fontSize=9, leading=12.6, leftIndent=12, bulletIndent=2, spaceAfter=2),
        "cel": ParagraphStyle("cel", fontName=F, fontSize=7.8, leading=10.6),
        "celc": ParagraphStyle("celc", fontName=F, fontSize=7.8, leading=10.6, alignment=1),
        "cab": ParagraphStyle("cab", fontName=FB, fontSize=7.8, leading=10.2, textColor=colors.white, alignment=1),
        "nota": ParagraphStyle("n", fontName=F, fontSize=7.4, leading=9.6, textColor=cinza, spaceAfter=3),
    }
    return S, F, verde, cinza


def gerar_pdf(res: dict, laudo: pd.DataFrame, comparacao: pd.DataFrame | None = None,
              alertas_historico: list[str] | None = None, titulo_extra: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    S, F, verde, cinza = _estilos()
    LARGURA = 17.4  # cm úteis no A4 com margens de 1,8 cm; todas as tabelas ocupam a largura inteira
    linha = colors.HexColor("#C9D3CB")

    def P(t, e="cel"):
        return Paragraph(_formula(str(t)), S[e])

    def tabela(dados, larguras):
        corpo = [[P(c, "cab") for c in dados[0]]] + [[P(c, "celc" if j else "cel") for j, c in enumerate(l)]
                                                    for l in dados[1:]]
        t = Table(corpo, colWidths=[w * cm for w in larguras], repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), verde), ("GRID", (0, 0), (-1, -1), 0.4, linha),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
        return t

    plano, ano = res["plano"], res["safra"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.5 * cm,
                            bottomMargin=1.6 * cm, title=f"Plano de adubação {ano}/{ano + 1}")

    def rodape(canv, d):
        canv.saveState()
        canv.setFont(F, 7.2)
        canv.setFillColor(cinza)
        canv.drawString(1.8 * cm, 1.0 * cm, "Sugestão técnica gerada a partir do laudo; validar doses com o "
                                           "responsável técnico (eng. agrônomo).")
        canv.drawRightString(A4[0] - 1.8 * cm, 1.0 * cm, f"Página {d.page}")
        canv.restoreState()

    st = [Paragraph(f"Plano de adubação · safra {ano}/{str(ano + 1)[-2:]}", S["titulo"]),
          Paragraph(f"{NIVEIS[res['nivel']]}{' · pastagens só com cloreto de potássio' if res['so_kcl'] else ''} · "
                    f"gerado em {date.today().strftime('%d/%m/%Y')}{' · ' + titulo_extra if titulo_extra else ''}",
                    S["sub"]), Spacer(1, 6)]

    st.append(Paragraph("1. O que comprar", S["h1"]))
    dados = [["Produto", "Total", "Embalagens", "Custo"]]
    for c in res["compras"].to_dict("records"):
        dados.append([f"<b>{c['produto']}</b>", f"{br_kg(c['kg'])} kg", f"{c['embalagens']} × {c['embalagem']}",
                      f"R$ {br(c['custo'], 2)}" if c["custo"] else "—"])
    st.append(tabela(dados, [6.5, 3, 4.5, 3.4]))

    st.append(Paragraph("2. Quanto aplicar em cada talhão", S["h1"]))
    pct = br(res.get("p2o5_pct", 26), 1 if res.get("p2o5_pct", 26) % 1 else 0)
    dados = [["Sigla", "Cultura", "Área", f"Superfosfato simples ({pct}% P₂O₅)", "25-00-25 no ano", "Cloreto de K",
              "Calcário"]]
    for p in plano.to_dict("records"):
        ssp = (f"{br(p['ssp_kg'])} kg<br/>{br(p['ssp_g_m2'])} g/m²" +
               (f"<br/>a cada {p['ssp_a_cada_anos']} anos" if p["ssp_a_cada_anos"] > 1 else "")) if p["ssp_kg_ha"] else "—"
        npk = (f"{br(p['npk_kg'])} kg<br/>{p['parcelas']} × {br(p['npk_g_m2'] / max(p['parcelas'], 1))} g/m²"
               if p["npk_kg_ha"] else "—")
        kcl = f"{br(p['kcl_kg'])} kg<br/>{br(p['kcl_g_m2'])} g/m²" if p["kcl_kg_ha"] else "—"
        cal = f"{br(p['calcario_kg'])} kg<br/>{p['calcario_tipo']}" if p["calcario_t_ha"] else "não"
        dados.append([f"<b>{p['sigla']}</b>", p["cultura_rotulo"], f"{br(p['area_m2'])} m²", ssp, npk, kcl, cal])
    st.append(tabela(dados, [1.4, 3.6, 1.9, 2.8, 3.1, 2.2, 2.4]))
    st.append(P("Doses por hectare: " + "; ".join(
        f"{p['sigla']} — superfosfato {br(p['ssp_kg_ha'])} ({br(p['P2O5_kg_ha'])} de P₂O₅), "
        f"25-00-25 {br(p['npk_kg_ha'])}"
        + (f", cloreto de K {br(p['kcl_kg_ha'])}" if p["kcl_kg_ha"] else "") + " kg/ha"
        for p in plano.to_dict("records")) + f". Superfosfato com {pct}% de P₂O₅: kg de adubo = kg de P₂O₅ ÷ "
        f"{br(res.get('p2o5_pct', 26) / 100, 2)}. À mão: g/m² = kg/ha ÷ 10.", "nota"))

    st.append(Paragraph("3. O que o laudo mostra", S["h1"]))
    dados = [["Sigla", "Laudo", "pH", "P (mg/dm³)", "K (mg/dm³)", "V% (meta)", "Ca/Mg", "Calagem"]]
    for p in plano.to_dict("records"):
        dados.append([f"<b>{p['sigla']}</b>", pd.to_datetime(p["data_laudo"]).strftime("%d/%m/%Y") if p["data_laudo"] else "—",
                      f"{br(p['pH'], 2)}<br/>{p['classe_pH'] or ''}", f"{br(p['P'])}<br/>{ROT[p['classe_P']]}",
                      f"{br(p['K_mg'])}<br/>{ROT[p['classe_K']]}", f"{br(p['V'])} ({br(p['V_meta'])})",
                      br(p["Ca_Mg"], 2), "não precisa" if not p["calcario_t_ha"] else f"{br(p['calcario_t_ha'], 2)} t/ha"])
    st.append(tabela(dados, [1.4, 2.0, 2.6, 2.2, 2.2, 2.2, 1.6, 3.2]))

    cal = res["calendario"]
    st.append(KeepTogether([Paragraph(f"4. Calendário da safra {ano}/{str(ano + 1)[-2:]}", S["h1"]),
                            tabela([["Sigla"] + list(cal.columns[1:])]
                                   + [[f"<b>{l[0]}</b>", *l[1:]] for l in cal.astype(str).values.tolist()],
                                   [2.6] + [(LARGURA - 2.6) / (cal.shape[1] - 1)] * (cal.shape[1] - 1))]))

    coq = res["coqueiros"]
    if not coq.empty:
        dados = [["Sigla", "Coqueiros", f"Superfosfato ({pct}%)", "Bórax", "25-00-25 no ano", "Total no talhão"]]
        for c in coq.to_dict("records"):
            dados.append([f"<b>{c['sigla']}</b>", f"{c['n_plantas']} ({c['idade']} anos)", f"{br(c['ssp_kg'], 1)} kg/planta",
                          f"{br(c['borax_g'])} g/planta" if c["borax_g"] else "—",
                          f"{br(c['npk_kg'], 1)} kg/planta ({c['parcelas']} × {br(c['kg_parcela'], 1)})",
                          f"{br(c['ssp_total_kg'])} kg SS + {br(c['npk_total_kg'])} kg 25-00-25"])
        st.append(KeepTogether([Paragraph("5. Coqueiros (adubação na coroa, de 0,5 a 2 m do tronco)", S["h1"]),
                                tabela(dados, [1.6, 2.6, 2.6, 2.0, 4.0, 4.6]),
                                Paragraph("Antes: coroamento de 2 a 2,5 m de raio, tirando capim e mato. Depois: "
                                          "misturar de leve e cobrir com palhada.", S["nota"])]))

    if comparacao is not None and not comparacao.empty:
        dados = [["Sigla", "Atributo", "Antes", "Agora", "Variação"]]
        casas = {rot: nd for _, rot, nd in CAMPOS_COMPARAR}
        dados += [[f"<b>{r['sigla']}</b>", r["atributo"], *(br(r[c], casas.get(r["atributo"], 2))
                                                             for c in ("antes", "agora", "variação"))]
                  for r in comparacao.to_dict("records")]
        st.append(KeepTogether([Paragraph("6. Comparação com a análise anterior", S["h1"]),
                                tabela(dados, [2, 4.4, 3.7, 3.7, 3.6])]))

    st.append(Paragraph("Alertas e observações", S["h1"]))
    for a in [*(alertas_historico or []), *res["alertas"]] or ["Nenhum alerta."]:
        st.append(Paragraph(a, S["bullet"], bulletText="•"))
    for cult, nota in sorted({(p["cultura_rotulo"], p["nota"]) for p in plano.to_dict("records")}):
        st.append(Paragraph(f"<b>{cult}:</b> {nota}", S["bullet"], bulletText="•"))
    st.append(Paragraph(
        "Referências: Freire et al. (2013), Manual de calagem e adubação do RJ; Cantarutti et al. (1999), "
        "5ª Aproximação (pastagens); Pereira et al. (2016), Embrapa CT 79 (BRS Capiaçu); Embrapa Tabuleiros "
        "Costeiros, A cultura do coqueiro.", S["nota"]))
    doc.build(st, onFirstPage=rodape, onLaterPages=rodape)
    return buf.getvalue()
