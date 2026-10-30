"""Testes de ponta a ponta: laudo PDF sintético no formato PESAGRO-RIO, planilha de outro laboratório,
plano de adubação, histórico, exportação e cópia do histórico.

Rodar:  python -m pytest -q
Com um laudo PDF real, defina LAUDO_PESAGRO=<caminho do PDF> para rodar também test_pdf_real.
"""
import io
import os
from pathlib import Path

import pandas as pd
import pytest

from laudo_sintetico import laudo_pdf
from solo.exportar import gerar_pdf, gerar_xlsx
from solo.historico import Historico, comparar
from solo.leitor import ModeloColunas, ler_laudo
from solo.motor import recomendar

TALHOES = {"T15": {"cultura": "capiacu", "area_m2": 5000}, "T815": {"cultura": "braquiaria", "area_m2": 12000},
           "T17": {"cultura": "braquiaria", "area_m2": 10000},
           "T45": {"cultura": "braquiaria_formacao", "area_m2": 3200}}


def texto_pdf(conteudo: bytes) -> str:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)


@pytest.fixture(scope="session")
def modelo(tmp_path_factory):
    return ModeloColunas(cache=tmp_path_factory.mktemp("m") / "modelo.joblib")


@pytest.fixture(scope="session")
def leitura_pdf(modelo):
    return ler_laudo(laudo_pdf(), "laudo_sintetico.pdf", modelo)


def test_pdf_formato_pesagro(leitura_pdf):
    d = leitura_pdf.dados.set_index("sigla")
    assert list(d.index) == ["T15", "T815", "T17", "T45"]
    assert leitura_pdf.meta["data_laudo"] == "2026-09-23"
    assert d.loc["T15", "P"] == 4 and d.loc["T15", "K_mg"] == 31 and d.loc["T15", "K_cmolc"] == 0.08
    assert d.loc["T45", "pH_agua"] == 6.10 and d.loc["T815", "V"] == 56 and d.loc["T17", "T"] == 6.68
    conf = leitura_pdf.colunas.loc[leitura_pdf.colunas["campo"] != "ignorar", "confianca"]
    assert len(conf) == 18 and conf.mean() > 0.9


@pytest.mark.skipif(not os.environ.get("LAUDO_PESAGRO"), reason="defina LAUDO_PESAGRO com o caminho de um laudo PDF")
def test_pdf_real(modelo):
    pdf = Path(os.environ["LAUDO_PESAGRO"])
    r = ler_laudo(pdf.read_bytes(), pdf.name, modelo)
    conf = r.colunas.loc[r.colunas["campo"] != "ignorar", "confianca"]
    assert len(r.dados) >= 1 and r.meta["data_laudo"] and len(conf) >= 15 and conf.mean() > 0.9
    assert r.dados["P"].notna().all() and r.dados["K_mg"].notna().all()


def test_planilha_com_outros_nomes(modelo):
    """Colunas embaralhadas, nomes diferentes, MO em dag/kg, K só em mg/dm³ e sem SB/T/V."""
    df = pd.DataFrame({
        "Fósforo (mg/dm3)": [4, 22], "Talhão": ["T-15", "T 815"], "pH H2O": ["5,0", "5,1"],
        "Cálcio (cmolc/dm3)": [2.92, 2.09], "Magnésio (cmolc/dm3)": [2.73, 1.10], "Potássio (mg/dm3)": [31, 54],
        "Matéria orgânica (dag/kg)": [3.38, 2.67], "Al (cmolc/dm3)": [0.2, 0.2], "H+Al (cmolc/dm3)": [2.6, 2.7],
    })
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    r = ler_laudo(buf.getvalue(), "outro_lab.xlsx", modelo)
    d = r.dados.set_index("sigla")
    assert list(d.index) == ["T15", "T815"]
    assert d.loc["T15", "P"] == 4 and d.loc["T815", "K_mg"] == 54
    assert abs(d.loc["T15", "MO_g_kg"] - 33.8) < 0.01           # dag/kg -> g/kg
    assert abs(d.loc["T15", "K_cmolc"] - 31 / 391) < 0.001       # derivado
    assert abs(d.loc["T15", "V"] - round(100 * (2.92 + 2.73 + 0.079) / (2.92 + 2.73 + 0.079 + 2.6))) <= 1


def _compra(res, inicio):
    c = res["compras"]
    return c[c["produto"].str.startswith(inicio)].iloc[0]


def test_plano_economico_com_superfosfato_comum(leitura_pdf):
    """Superfosfato simples comum (18% P2O5): 562 kg (12 sacos) e 644 kg de 25-00-25 (13 sacos)."""
    atual = leitura_pdf.dados.assign(data_laudo="2026-09-23")
    res = recomendar(atual, TALHOES, nivel="economico", p2o5_pct=18)
    p = res["plano"].set_index("sigla")
    assert (round(p.loc["T15", "ssp_kg_ha"]), p.loc["T15", "npk_kg_ha"]) == (600, 500)
    assert (round(p.loc["T815", "ssp_kg_ha"]), p.loc["T815", "npk_kg_ha"]) == (0, 150)
    assert (round(p.loc["T17", "ssp_kg_ha"]), p.loc["T17", "npk_kg_ha"]) == (150, 150)
    assert (round(p.loc["T45", "ssp_kg_ha"]), p.loc["T45", "npk_kg_ha"]) == (350, 200)
    ssp = _compra(res, "Superfosfato simples (18%")
    assert round(ssp["kg"]) == 562 and ssp["embalagens"] == 12
    npk = _compra(res, "25-00-25")
    assert round(npk["kg"]) == 644 and npk["embalagens"] == 13
    assert p["calcario_t_ha"].sum() == 0
    assert gerar_pdf(res, atual)[:4] == b"%PDF"
    assert gerar_xlsx(res, atual)[:2] == b"PK"


def test_superfosfato_reativo_26_por_cento(leitura_pdf):
    """Padrão do painel: superfosfato simples reativo com 26% de P2O5; a dose de P2O5 por hectare não muda."""
    atual = leitura_pdf.dados.assign(data_laudo="2026-09-23")
    talhoes = {**TALHOES, "T15": {**TALHOES["T15"], "coqueiros_n": 10, "coqueiros_idade": 9}}
    res = recomendar(atual, talhoes, nivel="economico")
    assert res["p2o5_pct"] == 26
    p = res["plano"].set_index("sigla")
    assert dict(p["P2O5_kg_ha"]) == {"T15": 108, "T815": 0, "T17": 27, "T45": 63}
    assert abs(p.loc["T15", "ssp_kg_ha"] - 108 / 0.26) < 1e-9
    assert res["coqueiros"].iloc[0]["ssp_kg"] == 1.2            # 300 g de P2O5 / 0,26 = 1,15 kg/planta
    ssp = _compra(res, "Superfosfato simples (26% P₂O₅)")
    esperado = 108 / 0.26 * 0.5 + 27 / 0.26 * 1.0 + 63 / 0.26 * 0.32 + 1.2 * 10
    assert abs(ssp["kg"] - round(esperado, 1)) < 0.05 and ssp["embalagens"] == 9
    texto = texto_pdf(gerar_pdf(res, atual))
    assert "Superfosfato simples (26%" in texto and "₂" not in texto   # P2O5 em <sub>, não em glifo que falta
    composto = recomendar(atual, TALHOES, nivel="economico", p2o5_pct=14)  # ex.: 04-14-08
    assert abs(composto["plano"].set_index("sigla").loc["T15", "ssp_kg_ha"] - 108 / 0.14) < 1e-9
    with pytest.raises(ValueError):
        recomendar(atual, TALHOES, p2o5_pct=0)


def test_historico_e_recomposicao(tmp_path, leitura_pdf):
    h = Historico(tmp_path)
    h.registrar(leitura_pdf.dados, {"data_laudo": "2026-09-23", "laboratorio": "PESAGRO-RIO"})
    h.salvar_talhoes(TALHOES)
    nova = leitura_pdf.dados[leitura_pdf.dados["sigla"] == "T15"].copy()
    nova.loc[:, ["P", "K_mg"]] = [16, 22]                         # P corrigido, K caindo
    casamento = h.casar(["T15", "X99"]).set_index("sigla")
    assert casamento.loc["T15", "situacao"] == "conhecida" and casamento.loc["X99", "situacao"] == "nova"
    _, alertas = comparar(nova.iloc[0], h.anterior("T15", "2027-09-20"))
    assert any("fósforo subiu" in a for a in alertas) and any("potássio caiu" in a for a in alertas)
    h.registrar(nova, {"data_laudo": "2027-09-20", "laboratorio": "PESAGRO-RIO"})
    atuais = h.atuais().set_index("sigla")
    assert atuais.loc["T15", "data_laudo"] == "2027-09-20" and atuais.loc["T17", "data_laudo"] == "2026-09-23"
    res = recomendar(h.atuais(), h.talhoes(), nivel="economico")
    assert res["plano"].set_index("sigla").loc["T15", "P2O5_kg_ha"] == 36   # P médio: só manutenção


def test_exportacao_leva_alertas_e_total(leitura_pdf):
    """O PDF e o Excel trazem os alertas do histórico, e o total da aba Compras fica na linha do rótulo."""
    from openpyxl import load_workbook

    atual = leitura_pdf.dados.assign(data_laudo="2026-09-23")
    res = recomendar(atual, TALHOES, nivel="economico")
    alerta = "T15: o potássio caiu 38% (50 → 31 mg/dm³)."
    texto = texto_pdf(gerar_pdf(res, atual, None, [alerta]))
    assert "potássio caiu 38%" in texto and "Nenhum alerta" not in texto
    wb = load_workbook(io.BytesIO(gerar_xlsx(res, atual, None, [alerta])))
    assert wb["Alertas"]["A2"].value == alerta
    ws = wb["Compras"]
    assert ws.cell(ws.max_row, 5).value == "Total" and ws.cell(ws.max_row, 6).value == f"=SUM(F2:F{ws.max_row - 1})"


def test_copia_do_historico_e_siglas_sem_rapidfuzz(tmp_path, leitura_pdf, monkeypatch):
    """O zip leva o histórico entre computadores; sem rapidfuzz (navegador) o difflib casa as siglas."""
    import builtins

    origem = Historico(tmp_path / "pc")
    origem.registrar(leitura_pdf.dados, {"data_laudo": "2026-09-23", "laboratorio": "PESAGRO-RIO"})
    origem.salvar_talhoes(TALHOES)
    destino = Historico(tmp_path / "outro")
    destino.importar_zip(origem.exportar_zip())
    assert destino.analises().equals(origem.analises()) and destino.talhoes() == origem.talhoes()
    with pytest.raises(ValueError):
        destino.importar_zip(b"isto nao e zip")

    importar = builtins.__import__
    monkeypatch.setattr(builtins, "__import__", lambda nome, *a, **k: (_ for _ in ()).throw(ImportError(nome))
                        if nome.startswith("rapidfuzz") else importar(nome, *a, **k))
    c = destino.casar(["T15", "T-15", "X99"]).set_index("sigla")
    assert c.loc["T15", "situacao"] == "conhecida" and c.loc["X99", "situacao"] == "nova"
    assert c.loc["T-15", "situacao"] == "parecida" and c.loc["T-15", "casar_com"] == "T15"
